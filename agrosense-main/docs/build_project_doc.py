from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AgroSense_Project_Documentation.docx"

GREEN = "0E9C6A"
DARK_GREEN = "073B35"
INK = "17323D"
BLUE = "2E74B5"
MUTED = "6F818A"
PALE_GREEN = "E8F7EE"
PALE_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "D9E4E7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_text(doc, text, style=None, bold=False, color=INK, size=11, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375 + (0.25 * level))
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(" " + text)
    set_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r)
    return p


def add_code(doc, lines):
    for line in lines.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F4F7F8")
        p_pr.append(shd)
        r = p.add_run(line if line else " ")
        set_font(r, name="Courier New", size=9, color=DARK_GREEN)


def add_callout(doc, label, text, fill=PALE_GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    set_font(r, size=10.5, color=DARK_GREEN, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=9.5, color=DARK_GREEN, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_font(r, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_GREEN, bold=True)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [(1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK_GREEN, 10, 5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section):
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("AGROSENSE  |  PROJECT REFERENCE GUIDE")
    set_font(r, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("AgroSense · ")
    set_font(r, size=8.5, color=MUTED)
    add_field(fp, "PAGE")


def build():
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    # Cover/title block: editorial-cover override for a technical reference guide.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROJECT REFERENCE GUIDE")
    set_font(r, size=10, color=GREEN, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AgroSense")
    set_font(r, size=30, color=DARK_GREEN, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Smart irrigation monitoring and control platform")
    set_font(r, size=15, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Current implementation guide · Version 0.1.0 · 7 September 2026")
    set_font(r, size=10.5, color=MUTED)
    add_callout(doc, "Purpose", "This document is the single project reference for the current AgroSense product scope, architecture, setup, API contracts, device integration, and implementation status.")

    add_heading(doc, "Start here", 1)
    add_text(doc, "Run the backend and frontend in separate terminals. The backend can start without PostgreSQL or MQTT; it uses seeded in-memory telemetry until those services are configured.")
    add_text(doc, "Terminal 1 · backend", bold=True, color=DARK_GREEN, after=2)
    add_code(doc, "cd backend\ncp .env.example .env\nnpm install\nnpm run dev")
    add_text(doc, "Terminal 2 · frontend", bold=True, color=DARK_GREEN, before=8, after=2)
    add_code(doc, "cd frontend\ncp .env.example .env\nnpm install\nnpm run dev")
    add_text(doc, "Open the dashboard at http://localhost:5173. The Express API runs at http://localhost:4000.", after=8)
    add_callout(doc, "Production build", "Use npm run build in both backend and frontend directories. Start the compiled backend with npm start.", fill=LIGHT_GRAY)

    add_heading(doc, "Project at a glance", 1)
    add_table(doc, ["Area", "Current choice"], [
        ("Product", "IoT-enabled smart irrigation dashboard for farm monitoring and control"),
        ("Frontend", "React, TypeScript, Vite, local shadcn-style UI primitives, Lucide icons"),
        ("Backend", "Node.js, Express, TypeScript, Socket.IO, MQTT.js, Zod"),
        ("Persistence", "PostgreSQL through pg; in-memory fallback for development"),
        ("Device", "ESP32 sensor/controller sending telemetry over Wi-Fi"),
        ("Telemetry protocol", "MQTT topic agrosense/{device_id}/telemetry"),
        ("Live UI protocol", "Socket.IO events telemetry:update and dashboard:summary"),
        ("Current status", "Dashboard shell, live summary, validation, MQTT gateway, and PostgreSQL write-through are implemented"),
    ], [2200, 7160])

    add_heading(doc, "Requirements from the supplied structure document", 1)
    add_text(doc, "The attached structure document is treated as the product/design reference. Its page descriptions are requirements for the AgroSense application, not additional instructions that override the user's request.")
    sections = [
        ("Login", "User login, password visibility, remember-me option, forgot-password flow, and optional Google login treatment."),
        ("Dashboard", "Temperature, humidity, water level, rain status, flow rate, all-zone soil moisture, pump status, water usage, and system status."),
        ("Live Sensor Data", "Real-time readings, zone-wise moisture, live moisture trends, current sensor status, tank status, and rain status."),
        ("Zones", "Independent monitoring and management of Zone 1, Zone 2, and Zone 3, including crops, pump state, trends, comparison, and zone irrigation details."),
        ("System Control", "Automatic/manual modes, zone pump control, moisture thresholds, schedules, flow monitoring, and water-tank status."),
        ("Irrigation History", "Past events with date/time, zone, moisture before/after, duration, water used, pump status, filters, and charts."),
        ("Plant Analysis", "Plant health, growth stage, recommended moisture, environmental conditions, plant response, water fingerprint, and recommendations."),
        ("Alerts", "Low moisture, high temperature, low water, rainfall, pump, offline, system alerts, history, and notification settings."),
        ("Settings", "Device ID, Wi-Fi settings, thresholds, notifications, system preferences, data management, and save settings."),
    ]
    add_table(doc, ["Module", "Scope"], sections, [2100, 7260])

    add_heading(doc, "Architecture", 1)
    add_text(doc, "The ESP32 communicates over Wi-Fi. MQTT is the recommended device-to-backend transport because it is lightweight, reconnectable, and suitable for continuous sensor messages.")
    add_code(doc, "ESP32 sensors / pumps\n        | Wi-Fi + MQTT\n        v\nExpress backend\n  |-- MQTT ingestion and validation\n  |-- PostgreSQL write-through\n  |-- Alert and irrigation services\n  |-- Socket.IO live updates\n        v\nReact dashboard")
    add_heading(doc, "Repository structure", 2)
    add_code(doc, "backend/\n  src/config              Environment parsing\n  src/middleware          Express error handling\n  src/routes              Versioned API routes\n  src/services            Telemetry store and MQTT service\n  src/types               Shared telemetry types\n  src/validation          Zod schemas\n  prisma/schema.prisma    Reference application data model\nfrontend/\n  src/components/ui       Local shadcn-style primitives\n  src/App.tsx             Dashboard and navigation shell\n  src/api.ts              REST client\n  src/styles.css          AgroSense visual system and layout\n  components.json         shadcn component configuration")

    add_heading(doc, "Frontend implementation", 1)
    add_text(doc, "The frontend currently provides the AgroSense dashboard shell and uses local shadcn-style components so the UI is editable and consistent within the repository.")
    add_bullet(doc, "Button: navigation, panel actions, and action variants")
    add_bullet(doc, "Card: metric cards, zone cards, and dashboard panels")
    add_bullet(doc, "Badge: device, connection, and pump status")
    add_bullet(doc, "Avatar: admin user chip")
    add_bullet(doc, "Progress: zone soil-moisture levels")
    add_bullet(doc, "Separator: topbar structure")
    add_bullet(doc, "Lucide icons: consistent navigation and metric icon alignment")
    add_text(doc, "Implemented dashboard content includes connection status, five system metrics, soil moisture by zone, water-usage visualization, moisture trend, pump donut, system status, and a branded sustainability banner.")

    add_heading(doc, "Environment configuration", 1)
    add_table(doc, ["Variable", "Default / example", "Purpose"], [
        ("PORT", "4000", "Express API port"),
        ("CLIENT_ORIGIN", "http://localhost:5173", "Allowed frontend origin"),
        ("DATABASE_URL", "empty", "PostgreSQL connection string; empty enables in-memory fallback"),
        ("MQTT_ENABLED", "false", "Enable ESP32 MQTT ingestion"),
        ("MQTT_URL", "mqtt://localhost:1883", "MQTT broker address"),
        ("MQTT_USERNAME / MQTT_PASSWORD", "empty", "Optional broker credentials"),
        ("MQTT_TOPIC_PREFIX", "agrosense", "Topic namespace prefix"),
        ("VITE_API_URL", "http://localhost:4000/api/v1", "Frontend REST base URL"),
        ("VITE_SOCKET_URL", "http://localhost:4000", "Frontend Socket.IO origin"),
    ], [2600, 2450, 4310])

    add_heading(doc, "Telemetry contract", 1)
    add_text(doc, "The backend accepts a payload with a device ID, system readings, and one or more zone readings. The timestamp is optional; the backend assigns the current time when omitted.")
    add_code(doc, '{\n  "device_id": "AGROSENSE01",\n  "timestamp": "2026-09-07T10:30:00Z",\n  "system": {\n    "temperature": 29.5,\n    "humidity": 64,\n    "water_level": 78,\n    "rain": false,\n    "flow_rate": 1.8\n  },\n  "zones": [\n    { "zone_id": 1, "soil_moisture": 28, "pump_status": true },\n    { "zone_id": 2, "soil_moisture": 52, "pump_status": false },\n    { "zone_id": 3, "soil_moisture": 21, "pump_status": true }\n  ]\n}')

    add_heading(doc, "Express API", 1)
    add_table(doc, ["Method", "Endpoint", "Purpose"], [
        ("GET", "/api/v1/health", "Service health check"),
        ("GET", "/api/v1/dashboard/summary", "Latest dashboard summary; optional device_id query parameter"),
        ("GET", "/api/v1/devices/:deviceId/latest", "Latest raw telemetry for a device"),
        ("POST", "/api/v1/devices/:deviceId/telemetry", "Validated HTTP telemetry simulator/fallback endpoint"),
    ], [1000, 3300, 5060])
    add_text(doc, "Example telemetry test:", bold=True, color=DARK_GREEN, after=2)
    add_code(doc, "curl -X POST http://localhost:4000/api/v1/devices/AGROSENSE01/telemetry \\\n  -H 'content-type: application/json' \\\n  --data @telemetry.json")

    add_heading(doc, "ESP32 and MQTT integration", 1)
    add_text(doc, "Set MQTT_ENABLED=true in backend/.env when the broker is available. The backend subscribes to:")
    add_code(doc, "agrosense/{device_id}/telemetry")
    add_text(doc, "The ESP32 should publish validated JSON telemetry periodically and reconnect automatically after Wi-Fi loss. Future control commands should use separate command and acknowledgement topics:")
    add_code(doc, "agrosense/{device_id}/commands\nagrosense/{device_id}/commands/ack")
    add_callout(doc, "Safety requirement", "The ESP32 should enforce local safety rules for low tank level, missing flow, maximum pump duration, and conflicting pump commands. It must not depend entirely on cloud connectivity for physical safety.", fill="FFF6DF")

    add_heading(doc, "Current implementation status", 1)
    add_table(doc, ["Area", "Status", "Notes"], [
        ("Express backend", "Implemented", "TypeScript, validation, health, dashboard summary, telemetry endpoint"),
        ("MQTT gateway", "Implemented", "Disabled by default; validates incoming ESP32 payloads and broadcasts live updates"),
        ("PostgreSQL write-through", "Implemented", "Creates telemetry tables when DATABASE_URL is configured; falls back safely otherwise"),
        ("React dashboard", "Implemented", "Dashboard shell, live connection state, cards, charts, zones, status panels"),
        ("Authentication", "Planned", "Login UI is in the product scope; protected auth routes are not yet implemented"),
        ("Pump control", "Planned", "Safety-oriented API and ESP32 command acknowledgement are next"),
        ("History, alerts, plant analysis, settings", "Planned", "Navigation placeholders exist; domain modules remain to be built"),
        ("Real hardware validation", "Pending", "Use an MQTT simulator before connecting physical pumps"),
    ], [2350, 1500, 5510])

    add_heading(doc, "Recommended build sequence", 1)
    for item in [
        "Confirm local setup with the seeded dashboard and HTTP telemetry endpoint.",
        "Run an MQTT broker and replace HTTP simulator messages with ESP32 telemetry.",
        "Connect PostgreSQL and verify telemetry hydration after backend restart.",
        "Implement authentication and protected routes.",
        "Build zones and safe manual/automatic irrigation control.",
        "Add irrigation history and CSV export.",
        "Add the alert rules engine and notification preferences.",
        "Build plant analysis from rule-based insights, then consider predictive models after historical data exists.",
        "Test device failure, stale data, low tank, rain, abnormal flow, and Wi-Fi loss scenarios.",
        "Deploy with HTTPS, a managed PostgreSQL database, an MQTT broker, and application monitoring.",
    ]:
        add_number(doc, item)

    add_heading(doc, "Verification checklist", 1)
    for item in [
        "backend npm run typecheck passes",
        "backend npm run build passes",
        "frontend npm run typecheck passes",
        "frontend npm run build passes",
        "GET /api/v1/health returns status ok",
        "POST telemetry accepts valid readings and updates the dashboard summary",
        "Invalid telemetry returns HTTP 400",
        "MQTT remains disabled until explicitly configured",
        "PostgreSQL fallback keeps local development usable without a database",
        "Frontend connects through REST and Socket.IO",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Troubleshooting", 1)
    add_heading(doc, "Port already in use", 2)
    add_text(doc, "The backend uses port 4000 and the frontend uses port 5173 by default. Stop the existing process or set PORT / Vite server configuration before starting another instance.")
    add_heading(doc, "Dashboard shows offline", 2)
    add_text(doc, "Start the backend, confirm the API URL, and submit telemetry. The in-memory store marks a device offline when its latest reading is more than five minutes old.")
    add_heading(doc, "PostgreSQL is unavailable", 2)
    add_text(doc, "Leave DATABASE_URL empty for local UI work. The backend logs the fallback and continues with seeded in-memory telemetry.")
    add_heading(doc, "MQTT is unavailable", 2)
    add_text(doc, "Leave MQTT_ENABLED=false and use the HTTP telemetry endpoint while the broker or ESP32 is not ready.")

    add_heading(doc, "Source and scope note", 1)
    add_text(doc, "The functional scope and sample data in this guide are consolidated from the supplied AgroSense structure document and the implementation currently present in this repository. Product screenshots are treated as visual direction; safety behavior, authentication policy, hardware wiring, notification channels, and plant-analysis formulas still require product decisions before production release.")

    doc.core_properties.title = "AgroSense Project Documentation"
    doc.core_properties.subject = "Project requirements, architecture, setup, API, MQTT, and implementation status"
    doc.core_properties.author = "AgroSense"
    doc.core_properties.keywords = "AgroSense, ESP32, MQTT, Express, React, irrigation"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
